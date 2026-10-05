import datetime
import logfire
from openpyxl import load_workbook


def format_cell(value) -> str:
    """Converts a cell value to clean text (no trailing .0 on whole numbers, ISO dates)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    # openpyxl reads date-only cells as datetimes at midnight
    if isinstance(value, datetime.datetime) and value.time() == datetime.time(0, 0):
        return value.date().isoformat()
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value).strip()


def parse_excel(file_path: str):
    """
    Parse Excel workbooks (.xlsx, .xlsm) using openpyxl.
    every sheet is processed, and each row is rendered as "column: value" pairs so chunks keep their headers for RAG.
    """

    with logfire.span("Excel Parsing", filename=file_path):
        try:
            # read_only streams large files; data_only returns cached formula results instead of formulas
            workbook = load_workbook(file_path, read_only=True, data_only=True)

            sheets = []
            total_rows = 0
            try:
                for sheet in workbook.worksheets:
                    # 1. Read rows and drop fully empty ones
                    rows = [
                        [format_cell(cell) for cell in row]
                        for row in sheet.iter_rows(values_only=True)
                    ]
                    rows = [row for row in rows if any(row)]

                    if not rows:
                        continue

                    # 2. First non-empty row is the header; fill in blank column names
                    headers = [h or f"column_{i + 1}" for i, h in enumerate(rows[0])]

                    # 3. Render each row as "header: value" lines, skipping empty cells
                    records = [f"Sheet: {sheet.title}"]
                    for row_num, row in enumerate(rows[1:], start=1):
                        lines = [f"Row {row_num}"]
                        for i, value in enumerate(row):
                            if not value:
                                continue
                            header = headers[i] if i < len(headers) else f"column_{i + 1}"
                            lines.append(f"{header}: {value}")
                        records.append("\n".join(lines))

                    total_rows += len(rows) - 1
                    sheets.append("\n\n".join(records))
            finally:
                # read_only workbooks keep the file handle open until closed
                workbook.close()

            full_text = "\n\n".join(sheets)

            if not full_text.strip():
                logfire.warning(f"Excel parser returned empty text for {file_path}")

            else:
                logfire.info(f"Successfully parsed {len(sheets)} sheets, {total_rows} rows, {len(full_text)} characters")

            return full_text

        except Exception as e:
            logfire.error(f"Excel Parse Failed: {e}")
            raise e
